"""
Research Bundle Service
Orchestrates the full research flow: query analysis, report generation,
and automatic bundling into downloadable PDF and DOCX documents.
"""
import os
import json
import time
import shutil
import tempfile
import asyncio
import threading
from uuid import uuid4
from pathlib import Path
from typing import Dict, List, Optional, Any, Tuple
from datetime import datetime
from loguru import logger


RESEARCH_BUNDLE_DIR = Path(tempfile.gettempdir()) / "amaniquery_research"
RESEARCH_BUNDLE_DIR.mkdir(parents=True, exist_ok=True)

# Clean bundles older than 1 hour periodically
_MAX_BUNDLE_AGE = 3600


class ResearchBundle:
    """Represents a completed research session with generated documents."""

    def __init__(self, bundle_id: str, query: str):
        self.bundle_id = bundle_id
        self.query = query
        self.created_at = time.time()
        self.analysis: Optional[Dict[str, Any]] = None
        self.report_content: Optional[str] = None
        self.report_metadata: Optional[Dict[str, Any]] = None
        self.pdf_path: Optional[str] = None
        self.docx_path: Optional[str] = None
        self.error: Optional[str] = None
        self.status: str = "pending"  # pending | running | completed | failed

    @property
    def age(self) -> float:
        return time.time() - self.created_at

    def to_dict(self) -> Dict[str, Any]:
        base = {
            "bundle_id": self.bundle_id,
            "query": self.query,
            "created_at": datetime.fromtimestamp(self.created_at).isoformat(),
            "status": self.status,
            "analysis_summary": None,
            "has_pdf": self.pdf_path is not None and os.path.exists(self.pdf_path),
            "has_docx": self.docx_path is not None and os.path.exists(self.docx_path),
            "download_urls": {},
        }
        if self.status == "completed":
            if self.analysis:
                raw = self.analysis
                analysis = raw.get("analysis", {})
                base["analysis_summary"] = (
                    analysis.get("query_interpretation", "")[:200]
                    if isinstance(analysis, dict)
                    else str(analysis)[:200]
                )
                base["sources"] = raw.get("sources", [])
                base["analysis"] = raw
            if base["has_pdf"]:
                base["download_urls"]["pdf"] = f"/research/download/{self.bundle_id}/pdf"
            if base["has_docx"]:
                base["download_urls"]["docx"] = f"/research/download/{self.bundle_id}/docx"
        if self.error:
            base["error"] = self.error
        return base


class ResearchBundleService:
    """
    Orchestrates full research flow:
      1. Analyze query via agentic (or legacy) research module
      2. Generate structured report
      3. Bundle into PDF + DOCX
      4. Return bundle with download URLs
    """

    def __init__(self, research_module=None, agentic_research_module=None,
                 report_generator=None, cache_manager=None):
        self.research_module = research_module
        self.agentic_research_module = agentic_research_module
        self.report_generator = report_generator
        self.cache_manager = cache_manager
        self._bundles: Dict[str, ResearchBundle] = {}

        self._cleaner_thread = threading.Thread(target=self._cleanup_loop, daemon=True)
        self._cleaner_thread.start()

    def _get_active_module(self):
        if self.agentic_research_module is not None:
            return self.agentic_research_module
        return self.research_module

    def _cleanup_loop(self):
        while True:
            time.sleep(120)
            now = time.time()
            stale = [bid for bid, b in self._bundles.items()
                     if now - b.created_at > _MAX_BUNDLE_AGE]
            for bid in stale:
                bundle = self._bundles.pop(bid, None)
                if bundle:
                    self._cleanup_bundle_files(bundle)
                    logger.info(f"Cleaned up stale research bundle: {bid}")

    def _cleanup_bundle_files(self, bundle: ResearchBundle):
        for path in [bundle.pdf_path, bundle.docx_path]:
            if path and os.path.exists(path):
                try:
                    os.remove(path)
                except Exception:
                    pass

    async def conduct_research(self, query: str, context: Optional[Dict] = None,
                               session_id: Optional[str] = None,
                               generate_pdf: bool = True,
                               generate_docx: bool = True) -> Dict[str, Any]:
        """
        Full research flow: analyze → report → bundle documents.
        Returns bundle metadata with download URLs.
        """
        bundle_id = uuid4().hex[:12]
        bundle = ResearchBundle(bundle_id, query)
        self._bundles[bundle_id] = bundle

        try:
            bundle.status = "running"

            # Step 1: Analyze the query
            module = self._get_active_module()
            if module is None:
                raise RuntimeError("No research module available")

            logger.info(f"[ResearchBundle:{bundle_id}] Analyzing query...")
            if hasattr(module, 'analyze_legal_query') and asyncio.iscoroutinefunction(module.analyze_legal_query):
                analysis = await module.analyze_legal_query(query, context)
            else:
                analysis = module.analyze_legal_query(query, context)

            if "error" in analysis:
                raise RuntimeError(analysis["error"])

            bundle.analysis = analysis
            logger.info(f"[ResearchBundle:{bundle_id}] Analysis complete")

            # Step 2: Generate report
            report_content = None
            report_metadata = None
            if self.report_generator is not None:
                try:
                    logger.info(f"[ResearchBundle:{bundle_id}] Generating report...")
                    report = self.report_generator.generate_legal_query_report(analysis)
                    if "error" not in report:
                        report_content = report.get("content", "")
                        report_metadata = report.get("metadata", {})
                        bundle.report_content = report_content
                        bundle.report_metadata = report_metadata
                except Exception as e:
                    logger.warning(f"[ResearchBundle:{bundle_id}] Report generation failed: {e}")

            # Step 3: Generate downloadable documents
            if generate_pdf:
                try:
                    logger.info(f"[ResearchBundle:{bundle_id}] Generating PDF...")
                    pdf_path = self._generate_pdf(analysis, bundle_id, report_content)
                    bundle.pdf_path = pdf_path
                except Exception as e:
                    logger.warning(f"[ResearchBundle:{bundle_id}] PDF generation failed: {e}")

            if generate_docx:
                try:
                    logger.info(f"[ResearchBundle:{bundle_id}] Generating DOCX...")
                    docx_path = self._generate_docx(analysis, bundle_id, report_content)
                    bundle.docx_path = docx_path
                except Exception as e:
                    logger.warning(f"[ResearchBundle:{bundle_id}] DOCX generation failed: {e}")

            bundle.status = "completed"
            logger.info(f"[ResearchBundle:{bundle_id}] Research complete")

            # Cache result if cache_manager available
            if self.cache_manager and not analysis.get("error"):
                try:
                    self.cache_manager.set(query, analysis, ttl_type="default")
                except Exception:
                    pass

            return bundle.to_dict()

        except Exception as e:
            logger.error(f"[ResearchBundle:{bundle_id}] Research failed: {e}")
            bundle.status = "failed"
            bundle.error = str(e)
            return bundle.to_dict()

    def _generate_pdf(self, analysis: Dict[str, Any], bundle_id: str,
                      report_content: Optional[str] = None) -> str:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                            PageBreak, Table, TableStyle)
            from reportlab.lib.units import inch
        except ImportError:
            raise ImportError("reportlab required. Run: pip install reportlab")

        pdf_path = str(RESEARCH_BUNDLE_DIR / f"{bundle_id}.pdf")
        doc = SimpleDocTemplate(pdf_path, pagesize=A4)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'],
                                     fontSize=18, spaceAfter=20, alignment=1)
        heading_style = ParagraphStyle('CustomHeading', parent=styles['Heading2'],
                                       fontSize=14, spaceAfter=12,
                                       textColor=colors.HexColor("#1a365d"))
        sub_style = ParagraphStyle('CustomSub', parent=styles['Heading3'],
                                   fontSize=12, spaceAfter=8,
                                   textColor=colors.HexColor("#2d3748"))
        normal = styles['Normal']
        normal.fontSize = 10
        normal.leading = 14

        story = []

        # Title page
        story.append(Spacer(1, 2 * inch))
        story.append(Paragraph("RESEARCH REPORT", title_style))
        story.append(Spacer(1, 0.3 * inch))
        story.append(Paragraph("AmaniQuery Research Intelligence", styles['Heading2']))
        story.append(Spacer(1, 0.5 * inch))
        story.append(Paragraph(
            f"<b>Query:</b> {analysis.get('original_query', 'Research Query')}", normal))
        story.append(Spacer(1, 0.15 * inch))
        story.append(Paragraph(
            f"<b>Generated:</b> {analysis.get('research_timestamp', datetime.utcnow().isoformat())}",
            normal))
        story.append(Spacer(1, 0.15 * inch))
        story.append(Paragraph(
            f"<b>Confidence:</b> {analysis.get('report_confidence', 'N/A')}", normal))
        story.append(PageBreak())

        # Analysis sections
        raw_analysis = analysis.get("analysis", analysis)

        exec_summary = None
        if isinstance(raw_analysis, dict):
            exec_summary = raw_analysis.get("query_interpretation")
            applicable_laws = raw_analysis.get("applicable_laws", [])
            legal_analysis = raw_analysis.get("legal_analysis")
            practical = raw_analysis.get("practical_guidance", {})
            considerations = raw_analysis.get("additional_considerations", [])
        else:
            applicable_laws = []
            legal_analysis = None
            practical = {}
            considerations = []

        if exec_summary:
            story.append(Paragraph("QUERY ANALYSIS", heading_style))
            story.append(Paragraph(exec_summary, normal))
            story.append(Spacer(1, 0.2 * inch))

        if applicable_laws and isinstance(applicable_laws, list):
            story.append(Paragraph("APPLICABLE LAWS", heading_style))
            for law in applicable_laws:
                if isinstance(law, str):
                    story.append(Paragraph(f"• {law}", normal))
                elif isinstance(law, dict):
                    story.append(Paragraph(
                        f"<b>{law.get('law_name', 'Law')}</b>", sub_style))
                    if law.get('citation'):
                        story.append(Paragraph(
                            f"<i>Citation:</i> {law['citation']}", normal))
            story.append(Spacer(1, 0.2 * inch))

        if legal_analysis:
            story.append(Paragraph("LEGAL ANALYSIS", heading_style))
            story.append(Paragraph(legal_analysis, normal))
            story.append(Spacer(1, 0.2 * inch))

        if practical:
            story.append(Paragraph("PRACTICAL GUIDANCE", heading_style))
            steps = practical.get("steps", [])
            if isinstance(steps, list):
                for i, step in enumerate(steps, 1):
                    story.append(Paragraph(f"{i}. {step}", normal))
            story.append(Spacer(1, 0.2 * inch))

        if considerations and isinstance(considerations, list):
            story.append(Paragraph("ADDITIONAL CONSIDERATIONS", heading_style))
            for c in considerations:
                story.append(Paragraph(f"• {c}", normal))
            story.append(Spacer(1, 0.2 * inch))

        # Report content if available
        if report_content:
            story.append(PageBreak())
            story.append(Paragraph("GENERATED REPORT", heading_style))
            for para in report_content.split("\n\n"):
                para = para.strip()
                if para:
                    story.append(Paragraph(para, normal))
                    story.append(Spacer(1, 0.1 * inch))

        # Sources
        sources = analysis.get("sources", [])
        if sources:
            story.append(Spacer(1, 0.3 * inch))
            story.append(Paragraph("SOURCES", heading_style))
            for i, src in enumerate(sources, 1):
                title = src.get("title", src.get("source_name", f"Source {i}"))
                url = src.get("url", "")
                if url:
                    story.append(Paragraph(f"[{i}] {title} — {url}", normal))
                else:
                    story.append(Paragraph(f"[{i}] {title}", normal))

        # Disclaimer
        story.append(Spacer(1, 0.3 * inch))
        story.append(Paragraph(
            "<i>This report is for informational purposes only and does not "
            "constitute legal advice. Consult a qualified legal professional "
            "for advice specific to your situation.</i>", normal))

        doc.build(story)
        return pdf_path

    def _generate_docx(self, analysis: Dict[str, Any], bundle_id: str,
                       report_content: Optional[str] = None) -> str:
        try:
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            raise ImportError("python-docx required. Run: pip install python-docx")

        docx_path = str(RESEARCH_BUNDLE_DIR / f"{bundle_id}.docx")
        doc = Document()

        # Title
        title = doc.add_heading("RESEARCH REPORT", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle = doc.add_heading("AmaniQuery Research Intelligence", 1)
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph(f"Query: {analysis.get('original_query', 'Research Query')}")
        doc.add_paragraph(
            f"Generated: {analysis.get('research_timestamp', datetime.utcnow().isoformat())}")
        doc.add_paragraph(
            f"Confidence: {analysis.get('report_confidence', 'N/A')}")
        doc.add_page_break()

        raw_analysis = analysis.get("analysis", analysis)

        exec_summary = None
        applicable_laws = []
        legal_analysis = None
        practical = {}
        considerations = []

        if isinstance(raw_analysis, dict):
            exec_summary = raw_analysis.get("query_interpretation")
            applicable_laws = raw_analysis.get("applicable_laws", [])
            legal_analysis = raw_analysis.get("legal_analysis")
            practical = raw_analysis.get("practical_guidance", {})
            considerations = raw_analysis.get("additional_considerations", [])

        if exec_summary:
            doc.add_heading("Query Analysis", 1)
            doc.add_paragraph(exec_summary)

        if applicable_laws and isinstance(applicable_laws, list):
            doc.add_heading("Applicable Laws", 1)
            for law in applicable_laws:
                if isinstance(law, str):
                    doc.add_paragraph(law, style="List Bullet")
                elif isinstance(law, dict):
                    p = doc.add_paragraph()
                    run = p.add_run(law.get("law_name", "Law"))
                    run.bold = True
                    if law.get("citation"):
                        doc.add_paragraph(f"Citation: {law['citation']}")

        if legal_analysis:
            doc.add_heading("Legal Analysis", 1)
            doc.add_paragraph(legal_analysis)

        if practical:
            doc.add_heading("Practical Guidance", 1)
            steps = practical.get("steps", [])
            if isinstance(steps, list):
                for step in steps:
                    doc.add_paragraph(step, style="List Bullet")

        if considerations and isinstance(considerations, list):
            doc.add_heading("Additional Considerations", 1)
            for c in considerations:
                doc.add_paragraph(c, style="List Bullet")

        if report_content:
            doc.add_page_break()
            doc.add_heading("Generated Report", 1)
            for para in report_content.split("\n\n"):
                para = para.strip()
                if para:
                    doc.add_paragraph(para)

        sources = analysis.get("sources", [])
        if sources:
            doc.add_heading("Sources", 1)
            for i, src in enumerate(sources, 1):
                title = src.get("title", src.get("source_name", f"Source {i}"))
                url = src.get("url", "")
                if url:
                    doc.add_paragraph(f"[{i}] {title} — {url}")
                else:
                    doc.add_paragraph(f"[{i}] {title}")

        doc.add_paragraph(
            "\nThis report is for informational purposes only and does not "
            "constitute legal advice. Consult a qualified legal professional "
            "for advice specific to your situation."
        )

        doc.save(docx_path)
        return docx_path

    def get_bundle(self, bundle_id: str) -> Optional[ResearchBundle]:
        return self._bundles.get(bundle_id)

    def get_bundle_download_path(self, bundle_id: str, fmt: str) -> Optional[str]:
        bundle = self._bundles.get(bundle_id)
        if bundle is None:
            return None
        if fmt == "pdf":
            return bundle.pdf_path
        elif fmt == "docx":
            return bundle.docx_path
        return None
